"""
tests/test_export.py — Unit tests for export service (CSV, Excel, PDF).
"""

import sys
import os
import io

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest
from services.export_service import (
    export_csv,
    export_excel_single,
    export_excel_district,
    export_pdf_district,
    HEADER_NAMES,
)


SAMPLE_RECORDS = [
    {
        "id": 1,
        "place_id": "ChIJpid1",
        "name": "St. Mary's High School",
        "phone": "+91-9876543210",
        "website": "https://stmarys.edu.in",
        "email": "info@stmarys.edu.in",
        "instagram": "https://www.instagram.com/stmarysschool",
        "instagram_status": "Verified from official website",
        "facebook": None,
        "linkedin": None,
        "whatsapp": None,
        "address": "123, Church Road, Andheri West",
        "city": "Mumbai",
        "district": "Mumbai Suburban",
        "state": "Maharashtra",
        "postal_code": "400058",
        "latitude": 19.136,
        "longitude": 72.827,
        "maps_url": "https://maps.google.com/?cid=ChIJpid1",
        "rating": 4.3,
        "review_count": 150,
        "business_status": "OPERATIONAL",
        "contact_source": "Google Places",
        "verification_status": "Found on Google Places",
        "outreach_status": "New",
        "is_verified": False,
        "notes": "",
        "created_at": "2026-09-01T10:00:00",
        "last_checked": None,
    },
    {
        "id": 2,
        "place_id": "ChIJpid2",
        "name": "Delhi Public School",
        "phone": "+91-8765432109",
        "website": None,
        "email": None,
        "instagram": None,
        "instagram_status": "Needs manual verification",
        "facebook": None,
        "linkedin": None,
        "whatsapp": None,
        "address": "456, MG Road",
        "city": "Pune",
        "district": "Pune",
        "state": "Maharashtra",
        "postal_code": "411001",
        "latitude": 18.52,
        "longitude": 73.856,
        "maps_url": "https://maps.google.com/?cid=ChIJpid2",
        "rating": 4.1,
        "review_count": 80,
        "business_status": "OPERATIONAL",
        "contact_source": "Google Places",
        "verification_status": "Found on Google Places",
        "outreach_status": "Interested",
        "is_verified": True,
        "notes": "Spoke to principal",
        "created_at": "2026-09-02T12:00:00",
        "last_checked": "2026-09-03T09:00:00",
    },
]


class TestExportCSV:
    def test_returns_bytes(self):
        result = export_csv(SAMPLE_RECORDS)
        assert isinstance(result, bytes)

    def test_non_empty(self):
        result = export_csv(SAMPLE_RECORDS)
        assert len(result) > 0

    def test_contains_header(self):
        result = export_csv(SAMPLE_RECORDS).decode("utf-8-sig")
        assert "School Name" in result

    def test_contains_data(self):
        result = export_csv(SAMPLE_RECORDS).decode("utf-8-sig")
        assert "St. Mary" in result
        assert "Delhi Public School" in result

    def test_no_api_key_exposure(self):
        result = export_csv(SAMPLE_RECORDS).decode("utf-8-sig")
        assert "GOOGLE_MAPS_API_KEY" not in result
        assert "AIza" not in result  # typical API key prefix

    def test_empty_records(self):
        result = export_csv([])
        assert isinstance(result, bytes)
        header = result.decode("utf-8-sig").strip()
        assert "School Name" in header  # header still present


class TestExportExcelSingle:
    def test_returns_bytes(self):
        result = export_excel_single(SAMPLE_RECORDS)
        assert isinstance(result, bytes)
        # Excel magic bytes (PK zip header)
        assert result[:2] == b"PK"

    def test_non_empty(self):
        result = export_excel_single(SAMPLE_RECORDS)
        assert len(result) > 0

    def test_can_be_read_back(self):
        import openpyxl
        result = export_excel_single(SAMPLE_RECORDS)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        # First row should be headers
        headers = [cell.value for cell in ws[1]]
        assert "School Name" in headers
        assert "Phone" in headers

    def test_data_rows_present(self):
        import openpyxl
        result = export_excel_single(SAMPLE_RECORDS)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws.max_row >= 3  # 1 header + 2 data rows

    def test_empty_records(self):
        result = export_excel_single([])
        assert isinstance(result, bytes)


class TestExportExcelDistrict:
    def test_returns_bytes(self):
        result = export_excel_district(SAMPLE_RECORDS)
        assert isinstance(result, bytes)
        assert result[:2] == b"PK"

    def test_creates_district_sheets(self):
        import openpyxl
        result = export_excel_district(SAMPLE_RECORDS)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        sheet_names = wb.sheetnames
        # Records span Mumbai Suburban and Pune
        assert "Mumbai Suburban" in sheet_names or "Pune" in sheet_names

    def test_single_district(self):
        import openpyxl
        records = [r for r in SAMPLE_RECORDS if r["district"] == "Pune"]
        result = export_excel_district(records)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        assert "Pune" in wb.sheetnames


class TestExportPDFDistrict:
    def test_returns_bytes(self):
        result = export_pdf_district(SAMPLE_RECORDS, state="Maharashtra")
        assert isinstance(result, bytes)

    def test_is_pdf(self):
        result = export_pdf_district(SAMPLE_RECORDS, state="Maharashtra")
        # PDF magic bytes
        assert result[:4] == b"%PDF"

    def test_non_empty(self):
        result = export_pdf_district(SAMPLE_RECORDS, state="Maharashtra")
        assert len(result) > 1000  # Should be a meaningful PDF

    def test_empty_records(self):
        result = export_pdf_district([], state="Maharashtra")
        assert isinstance(result, bytes)
        assert result[:4] == b"%PDF"
